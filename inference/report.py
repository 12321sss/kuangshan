"""
矿山稳定性/安全性报告生成（S4）

基于数据库真实统计（时序变化、解译结果、传感器监测）生成 Word 报告。
python-docx 模板化，结论措辞规则化 —— 严格遵守口径红线：
  - 输出「潜在水土流失风险筛查」，不报实测侵蚀量
  - 「持续复绿」≠「生态修复成功」
  - 预测是趋势外推，带不确定性
"""
from __future__ import annotations

import io
from datetime import datetime

import json

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def _add_heading(doc: Document, text: str, level: int = 1) -> None:
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x0D, 0x94, 0x88)  # 品牌青绿


def _add_kv(doc: Document, label: str, value: str) -> None:
    p = doc.add_paragraph()
    r = p.add_run(f"{label}：")
    r.bold = True
    p.add_run(str(value))


def _add_note(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x86, 0x90, 0x9C)


def generate_report(data: dict) -> bytes:
    """生成报告，返回 docx 字节流。

    data 结构（由 Java 汇总后传入）：
      {
        "projectName": str,
        "generatedAt": str,
        "timeseries": [{"date", "totalAreaHa", "byClass": {...}}, ...],
        "summary": [{"classType", "totalAreaHa", "recordCount"}, ...],
        "forecast": {"insufficient", "forecast": [{"date","areaHa","lowerHa","upperHa"}]},
        "sensorOnline": int, "sensorAlarm": int,
      }
    """
    doc = Document()

    # ---- 标题 ----
    title = doc.add_heading("矿山遥感监测综合分析报告", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # ---- 一、矿区概况 ----
    _add_heading(doc, "一、矿区概况", level=1)
    _add_kv(doc, "监测项目", data.get("projectName", "-"))
    _add_kv(doc, "报告生成时间", data.get("generatedAt", datetime.now().strftime("%Y-%m-%d %H:%M")))
    _add_kv(doc, "数据来源", "Sentinel-2 遥感影像在线解译 + 传感器监测（全部来自数据库真实记录）")

    # ---- 二、地物时序变化 ----
    _add_heading(doc, "二、地物面积时序变化", level=1)
    ts = data.get("timeseries", [])
    if not ts:
        doc.add_paragraph("当前项目暂无解译记录。")
    else:
        table = doc.add_table(rows=1, cols=3)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        hdr[0].text = "影像日期"
        hdr[1].text = "合计面积(km²)"
        hdr[2].text = "较上期变化(km²)"
        for row in ts:
            cells = table.add_row().cells
            cells[0].text = str(row.get("date", "-"))
            cells[1].text = f"{row.get('totalAreaHa', 0) / 100:.2f}"
            chg = row.get("changeHa", 0)
            cells[2].text = f"{chg / 100:+.2f}"

    # ---- 三、地物构成 ----
    _add_heading(doc, "三、地物构成", level=1)
    summary = data.get("summary", [])
    if summary:
        table = doc.add_table(rows=1, cols=3)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        hdr[0].text = "地物类型"
        hdr[1].text = "累计面积(km²)"
        hdr[2].text = "记录数"
        for c in summary:
            cells = table.add_row().cells
            cells[0].text = str(c.get("classType", "-"))
            cells[1].text = f"{c.get('totalAreaHa', 0) / 100:.2f}"
            cells[2].text = str(c.get("recordCount", 0))

    # ---- 四、采坑趋势外推 ----
    _add_heading(doc, "四、采坑扩张趋势（外推）", level=1)
    forecast = data.get("forecast", {})
    if forecast.get("insufficient"):
        doc.add_paragraph("采坑类真实提取记录不足 2 期，无法做趋势外推。")
    else:
        table = doc.add_table(rows=1, cols=4)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        hdr[0].text = "外推日期"
        hdr[1].text = "面积中值(km²)"
        hdr[2].text = "区间下限(km²)"
        hdr[3].text = "区间上限(km²)"
        for f in forecast.get("forecast", []):
            cells = table.add_row().cells
            cells[0].text = str(f.get("date", "-"))
            cells[1].text = f"{f.get('areaHa', 0) / 100:.2f}"
            cells[2].text = f"{f.get('lowerHa', 0) / 100:.2f}"
            cells[3].text = f"{f.get('upperHa', 0) / 100:.2f}"
        _add_note(doc, "注：趋势外推非实测结果，实际扩张速率受采矿计划、地形与政策等因素影响。")

    # ---- 五、生态质量评价（RSEI）----
    rsei = data.get("rsei")
    if rsei and rsei.get("mean") is not None:
        _add_rsei_section(doc, rsei)

    # ---- 六、安全风险评估（多因子融合）----
    safety = data.get("safety")
    if safety and safety.get("available") and safety.get("score"):
        _add_safety_section(doc, safety)

    # ---- 七、传感器监测概况 ----
    _add_heading(doc, "七、传感器监测概况", level=1)
    _add_kv(doc, "在线设备数", str(data.get("sensorOnline", 0)))
    _add_kv(doc, "未处理告警数", str(data.get("sensorAlarm", 0)))

    _add_extraction_appendix(doc, data.get("extractions") or [])

    _add_extraction_appendix(doc, data.get("extractions") or [])

    # ---- 八、结论与口径说明 ----
    _add_heading(doc, "八、结论与口径说明", level=1)
    doc.add_paragraph("本报告输出为「潜在水土流失风险筛查」，非实测侵蚀量。")
    doc.add_paragraph("「持续复绿」仅表示植被覆盖度回升，不等于生态修复成功。")
    if safety and safety.get("available"):
        doc.add_paragraph(
            "安全风险评估输出为「潜在安全风险筛查」，不构成边坡稳定性评价，不给出安全系数；"
            "岩土参数、地下水与结构面条件需现场勘察，安全生产决策须以有资质单位的专项评估为准。"
        )
    _add_note(doc, "本报告由智慧矿山平台自动生成，数据可追溯到每次解译记录。")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


METHOD_LABEL = {"INDEX": "指数阈值法", "SUPERVISED": "监督分类", "FUSION": "分类融合"}
CLASS_LABEL = {
    "PIT": "露天采坑", "DUMP": "排土场", "TAILINGS": "尾矿设施", "BARE": "裸地扰动",
    "ROAD": "道路设施", "VEGETATION": "植被", "WATER": "水体",
    "BUILDING": "建筑", "MOUNTAIN": "山体",
}


def _add_extraction_appendix(doc: Document, extractions: list[dict]) -> None:
    """附录：解译成果明细 + 分类精度

    为什么单独做附录：正文的时序/构成是聚合口径，客户核对时想看"每条记录什么时候、
    什么方法、多少面积"，这些明细放正文会淹没结论，放附录正好。
    """
    _add_heading(doc, "九、附录：遥感解译成果明细", level=1)
    if not extractions:
        doc.add_paragraph("当前项目暂无解译记录。")
        return

    doc.add_paragraph(f"共 {len(extractions)} 条解译记录（按下表逐条列示，便于复核）。")
    table = doc.add_table(rows=1, cols=6)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    for i, t in enumerate(["影像日期", "地物类型", "解译方法", "面积(km²)", "覆盖率", "数据源"]):
        hdr[i].text = t
    for e in extractions[:120]:
        cells = table.add_row().cells
        cells[0].text = str(e.get("acquisitionDate") or "-")
        cells[1].text = CLASS_LABEL.get(str(e.get("classType")), str(e.get("classType") or "-"))
        cells[2].text = METHOD_LABEL.get(str(e.get("method") or "INDEX"), str(e.get("method") or "INDEX"))
        cells[3].text = f"{(e.get('extractAreaHa') or 0) / 100:.2f}"
        cells[4].text = f"{e.get('coveragePct') or 0:.1f}%"
        cells[5].text = str(e.get("satellite") or "-")
    if len(extractions) > 120:
        doc.add_paragraph(f"（仅列示前 120 条，其余 {len(extractions) - 120} 条见系统记录）")

    # 分类精度：取最近一条带 accuracyJson 的监督分类记录
    acc = None
    for e in extractions:
        raw = e.get("accuracyJson")
        if e.get("method") == "SUPERVISED" and raw:
            try:
                acc = json.loads(raw)
                break
            except Exception:  # noqa: BLE001
                continue
    if acc:
        _add_heading(doc, "9.1 监督分类精度", level=2)
        _add_kv(doc, "总体精度 OA", f"{acc.get('overallAccuracy', '-')}%")
        _add_kv(doc, "Kappa 系数", str(acc.get("kappa", "-")))
        _add_kv(doc, "样本/验证", f"{acc.get('sampleCount', '-')} / {acc.get('split', '-')}")
        _add_note(doc, "精度指标只对本次分析的影像与 AOI 有效，不代表其它矿区或其它时相。")

    _add_note(doc,
              "上表面积均为「潜在水土流失/扰动风险筛查」口径的影像解译结果，"
              "不是实测侵蚀量，也不构成储量或品位数据。")


METHOD_LABEL = {"INDEX": "指数阈值法", "SUPERVISED": "监督分类", "FUSION": "分类融合"}
CLASS_LABEL = {
    "PIT": "露天采坑", "DUMP": "排土场", "TAILINGS": "尾矿设施", "BARE": "裸地扰动",
    "ROAD": "道路设施", "VEGETATION": "植被", "WATER": "水体",
    "BUILDING": "建筑", "MOUNTAIN": "山体",
}


def _add_extraction_appendix(doc: Document, extractions: list[dict]) -> None:
    """附录：解译成果明细 + 分类精度

    为什么单独做附录：正文的时序/构成是聚合口径，客户核对时想看"每条记录什么时候、
    什么方法、多少面积"，这些明细放正文会淹没结论，放附录正好。
    """
    _add_heading(doc, "九、附录：遥感解译成果明细", level=1)
    if not extractions:
        doc.add_paragraph("当前项目暂无解译记录。")
        return

    doc.add_paragraph(f"共 {len(extractions)} 条解译记录（按下表逐条列示，便于复核）。")
    table = doc.add_table(rows=1, cols=6)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    for i, t in enumerate(["影像日期", "地物类型", "解译方法", "面积(km²)", "覆盖率", "数据源"]):
        hdr[i].text = t
    for e in extractions[:120]:
        cells = table.add_row().cells
        cells[0].text = str(e.get("acquisitionDate") or "-")
        cells[1].text = CLASS_LABEL.get(str(e.get("classType")), str(e.get("classType") or "-"))
        cells[2].text = METHOD_LABEL.get(str(e.get("method") or "INDEX"), str(e.get("method") or "INDEX"))
        cells[3].text = f"{(e.get('extractAreaHa') or 0) / 100:.2f}"
        cells[4].text = f"{e.get('coveragePct') or 0:.1f}%"
        cells[5].text = str(e.get("satellite") or "-")
    if len(extractions) > 120:
        doc.add_paragraph(f"（仅列示前 120 条，其余 {len(extractions) - 120} 条见系统记录）")

    # 分类精度：取最近一条带 accuracyJson 的监督分类记录
    acc = None
    for e in extractions:
        raw = e.get("accuracyJson")
        if e.get("method") == "SUPERVISED" and raw:
            try:
                acc = json.loads(raw)
                break
            except Exception:  # noqa: BLE001
                continue
    if acc:
        _add_heading(doc, "9.1 监督分类精度", level=2)
        _add_kv(doc, "总体精度 OA", f"{acc.get('overallAccuracy', '-')}%")
        _add_kv(doc, "Kappa 系数", str(acc.get("kappa", "-")))
        _add_kv(doc, "样本/验证", f"{acc.get('sampleCount', '-')} / {acc.get('split', '-')}")
        _add_note(doc, "精度指标只对本次分析的影像与 AOI 有效，不代表其它矿区或其它时相。")

    _add_note(doc,
              "上表面积均为「潜在水土流失/扰动风险筛查」口径的影像解译结果，"
              "不是实测侵蚀量，也不构成储量或品位数据。")


def _add_rsei_section(doc: Document, rsei: dict) -> None:
    """生态质量评价章节（遥感生态指数 RSEI）

    口径：三因子 RSEI（绿度/湿度/干度），不含热环境维度——
    Sentinel-2 无热红外波段，无法反演地表温度。必须在正文写明，避免客户误当完整 RSEI。
    """
    _add_heading(doc, "五、生态质量评价（RSEI）", level=1)
    _add_kv(doc, "RSEI 均值", str(rsei.get("mean")))
    if rsei.get("datetime"):
        _add_kv(doc, "影像日期", str(rsei.get("datetime"))[:10])
    if rsei.get("sceneId"):
        _add_kv(doc, "数据来源", str(rsei.get("sceneId")))

    levels = rsei.get("levels") or []
    if levels:
        doc.add_paragraph("生态质量分级占比（RSEI 0~1，越高越好）：")
        table = doc.add_table(rows=1, cols=3)
        table.style = "Light Grid Accent 1"
        head = table.rows[0].cells
        head[0].text = "等级"
        head[1].text = "RSEI 区间"
        head[2].text = "面积占比"
        for l in levels:
            cells = table.add_row().cells
            cells[0].text = str(l.get("label", ""))
            cells[1].text = str(l.get("range", ""))
            cells[2].text = f"{l.get('pct', 0)}%"

    # 主成分权重：说明「哪一项主导了本区生态状况」
    weights = rsei.get("pc1Weight") or []
    labels = rsei.get("weightLabels") or ["绿度 NDVI", "湿度 WET", "干度 NDBSI"]
    if weights and len(weights) == len(labels):
        doc.add_paragraph("第一主成分（PC1）各指标载荷：")
        for lab, w in zip(labels, weights):
            doc.add_paragraph(f"{lab}：{w}", style="List Bullet")
        if rsei.get("explainRatio"):
            _add_note(doc, f"PC1 解释方差比例 {rsei['explainRatio']}，"
                           "比例越高说明这三项指标越能代表本区生态状况的整体差异。")

    # 两期对比结论
    cmp_ = rsei.get("compare")
    if cmp_:
        doc.add_paragraph("与上一期对比：")
        _add_kv(doc, "前期 RSEI 均值", str(cmp_.get("earlyMean")))
        _add_kv(doc, "本期 RSEI 均值", str(cmp_.get("lateMean")))
        _add_kv(doc, "变化", str(cmp_.get("deltaMean")))
        _add_kv(doc, "结论", str(cmp_.get("verdict", "")))
        _add_kv(doc, "改善/稳定/退化占比",
                f"{cmp_.get('improvedPct')}% / {cmp_.get('stablePct')}% / {cmp_.get('degradedPct')}%")

    _add_note(doc, "口径说明：本 RSEI 为三因子版本（绿度 NDVI + 湿度 WET + 干度 NDBSI），"
                   "不含热环境维度——Sentinel-2 没有热红外波段，无法反演地表温度。"
                   "完整四因子 RSEI 需融合 Landsat 8/9 热红外数据，属后续增强项。"
                   "RSEI 反映生态质量相对状况，不能直接等同于生态修复成效。")


def _add_safety_section(doc: Document, safety: dict) -> None:
    """安全风险评估章节

    设计原则：不只给总分，必须把「每一维的实测数值 + 建议」写进报告，
    否则报告无法支撑现场核查，也无法在质询时说明结论来源。
    """
    score = safety["score"]
    _add_heading(doc, "六、开采安全风险评估", level=1)

    _add_kv(doc, "综合风险分", f"{score.get('total')} / 100（{score.get('level')}风险）")
    if score.get("safetyIndex") is not None:
        _add_kv(doc, "矿山安全指数", f"{score.get('safetyIndex')} / 100（{score.get('safetyLevel')}）")
    ahp = score.get("ahp") or {}
    if ahp:
        _add_note(doc, f"权重方法：{score.get('weightMethod', 'AHP 层次分析法')}；"
                       f"一致性检验 CR = {ahp.get('CR')}（CR < 0.1 视为可接受），"
                       f"λmax = {ahp.get('lambdaMax')}。")
    if safety.get("projectName"):
        _add_kv(doc, "评估项目", str(safety["projectName"]))
    if safety.get("recordCount"):
        _add_kv(doc, "参与评估记录数", str(safety["recordCount"]))

    # 五维评分表
    dims = score.get("dimensions") or []
    if dims:
        doc.add_paragraph("各维度得分（加权后构成综合风险分）：")
        table = doc.add_table(rows=1, cols=4)
        table.style = "Light Grid Accent 1"
        head = table.rows[0].cells
        head[0].text = "维度"
        head[1].text = "得分"
        head[2].text = "权重"
        head[3].text = "加权得分"
        for d in dims:
            cells = table.add_row().cells
            cells[0].text = str(d.get("label", ""))
            cells[1].text = str(d.get("score", ""))
            cells[2].text = f"{round(float(d.get('weight', 0)) * 100)}%"
            cells[3].text = str(d.get("weighted", ""))

    # 逐维度依据
    for d in dims:
        doc.add_paragraph(f"【{d.get('label')}】得分 {d.get('score')}，数据可信度 {d.get('confidence', 'medium')}")
        for e in (d.get("evidence") or []):
            doc.add_paragraph(str(e), style="List Bullet")
        if d.get("advice"):
            _add_note(doc, f"建议：{d['advice']}")

    # 越界结论
    enc = safety.get("encroachment") or {}
    if enc.get("available"):
        doc.add_paragraph("越界开采检测：")
        _add_kv(doc, "越界面积", f"{enc.get('areaHa')} ha")
        _add_kv(doc, "占扰动总面积", f"{enc.get('pctOfDisturbance')}%")
        if (enc.get("areaHa") or 0) > 0:
            _add_note(doc, "检测到扰动范围超出登记矿权边界，需先排除解译边界误差与边界日期不一致，"
                          "再核对采矿许可范围；确认越界属重大事项，应按管理规定上报。")
        else:
            _add_note(doc, "扰动范围未超出登记矿权边界（基于本次解译结果与当前矿权边界）。")

    # 热点清单
    hotspots = safety.get("hotspots") or []
    if hotspots:
        doc.add_paragraph(f"风险热点（前 {min(len(hotspots), 5)} 处，按风险分×面积排序）：")
        for i, h in enumerate(hotspots[:5], 1):
            doc.add_paragraph(
                f"{i}. 位置 {h.get('lng')}, {h.get('lat')}，面积 {h.get('areaHa')} ha，"
                f"平均风险分 {h.get('score')}（峰值 {h.get('maxScore')}）",
                style="List Number"
            )
    else:
        _add_note(doc, "本次未发现超过阈值的连片高风险区域。")

    _add_note(doc, str(safety.get("disclaimer") or ""))

